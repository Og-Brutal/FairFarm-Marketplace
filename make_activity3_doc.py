import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_activity3_docx():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette
    PRIMARY_COLOR = RGBColor(85, 84, 36)   # Olive #555424
    ACCENT_COLOR = RGBColor(133, 61, 37)   # Rust #853D25
    TEXT_COLOR = RGBColor(40, 40, 40)
    DARK_BG_HEX = "555424"
    LIGHT_BG_HEX = "F4F6F0"
    BORDER_HEX = "D1E2C4"

    # Document Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.paragraph_format.line_spacing = 1.15
    
    r_title = title_p.add_run("DEVELOPMENT ACTIVITY 3 - PARALLEL MODULE DEVELOPMENT & INTEGRATION")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(14)
    
    r_sub = sub_p.add_run("FairFarm - Farmer Marketplace with Price Transparency | Secure Software Development")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = ACCENT_COLOR

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    def add_key_value(label, val):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = TEXT_COLOR
        
        r_val = p.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10)
        r_val.font.color.rgb = TEXT_COLOR

    # SUBMISSION INFO
    add_section_header("Submission Overview & Project Information")
    add_key_value("Project Name", "FairFarm - Farmer Marketplace with Price Transparency")
    add_key_value("Group Members", "Abdul-Wahab (2026-SSD-01), Ahmad-Fareed (2026-SSD-02)")
    add_key_value("Public Website URL", "https://fair-farm-marketplace-p81iwnhzy-ogbrutal2825-8857s-projects.vercel.app/")
    add_key_value("Module 1", "Submit Crop Listing (SubmitListing.jsx) — Developed by Abdul-Wahab")
    add_key_value("Module 2", "Browse & Search Crop Listings (BrowseListings.jsx) — Developed by Ahmad-Fareed")

    # PART 1: SELECT TWO COMPLEMENTARY MODULES
    add_section_header("Part 1 — Selected Complementary Modules")
    add_key_value("Module 1 Definition", "Submit Crop Listing: Enables farmers to create and publish crop listings with detailed specifications (Crop type, variety, quantity, quality grade, minimum bidding price, district, description).")
    add_key_value("Module 2 Definition", "Browse & Search Listings: Enables buyers, officers, and farmers to discover produce listings through real-time keyword search, quality/status filters, price sorting, and interactive detail modals.")

    # PART 2: MEMBER 1 - MODULE 1
    add_section_header("Part 2 — Member 1: Develop Module 1 (Abdul-Wahab)")
    add_key_value("Responsible Member", "Abdul-Wahab")
    add_key_value("Primary Function", "Allows farmers to submit new agricultural produce for auction with complete quality and location metadata.")
    add_key_value("Input Validation Rules", "1) Mandatory fields required (crop, variety, quantity, quality, minPrice, district). 2) Variety length >= 2 chars. 3) Quantity & Min Price must be positive numbers (> 0). 4) Description length >= 10 chars if provided.")
    add_key_value("JavaScript Interaction Implemented", "Real-time client-side form validation, stateful session listing append (pushing newly created produce items to live session list), and animated toast feedback notification.")

    # PART 3: MEMBER 2 - MODULE 2
    add_section_header("Part 3 — Member 2: Develop Module 2 (Ahmad-Fareed)")
    add_key_value("Responsible Member", "Ahmad-Fareed")
    add_key_value("Primary Function", "Allows users to explore, filter, sort, and inspect crop listings across districts and categories.")
    add_key_value("Input Validation / Search Rules", "1) Live search matches produce name, variety, or district. 2) Filter parameters check crop type, quality tier, and bidding status (OPEN, ENDED, SOLD). 3) Empty result fallback view provides a 'Clear All Filters' reset handler.")
    add_key_value("JavaScript Interaction Implemented", "Dynamic array filtering and sorting via useMemo, removable active filter tag chips (removeFilter / clearAllFilters), and Framer Motion modal dialog for full produce inspection & bid simulation.")

    # PART 4: INTEGRATE BOTH MODULES
    add_section_header("Part 4 — Integration Checklist")
    add_key_value("Module 1 Status", "Accessible via route /submit-listing and integrated navbar link for logged-in Farmers.")
    add_key_value("Module 2 Status", "Accessible via route /browse-listings and default public navbar link.")
    add_key_value("Design Consistency", "Maintains FairFarm brand design tokens (Olive #555424, Rust #853D25, Sage #D1E2C4, Farm Cream #faf8f5) with responsive Tailwind layout and Framer Motion transitions.")

    # PART 5: TEST BOTH MODULES
    add_section_header("Part 5 — Testing Results Matrix")
    
    # Table for Test Cases
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    headers = ["Test Case", "Expected Result", "Actual Result"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], DARK_BG_HEX)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9.5)

    test_data = [
        ("Submit valid crop listing (Module 1)", "Action succeeds, toast feedback appears, card added to session list.", "Pass - Listing added to session sidebar immediately."),
        ("Leave required fields empty (Module 1)", "Submission prevented, inline validation errors highlighted.", "Pass - Form blocked submission and displayed red warning messages."),
        ("Search and filter listings (Module 2)", "Grid updates dynamically based on query and filter parameters.", "Pass - Results filtered instantly by keyword, quality, and status.")
    ]

    for row_idx, data in enumerate(test_data, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            if row_idx % 2 == 1:
                set_cell_background(row_cells[col_idx], LIGHT_BG_HEX)
            p = row_cells[col_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.size = Pt(9)
                p.runs[0].font.name = "Arial"

    add_key_value("Module 1 Issue Discovered", "Submitting form retained previous input data causing accidental duplicate submissions.")
    add_key_value("Module 2 Issue Discovered", "Empty search query returned broken grid layout when no produce cards matched.")
    add_key_value("Correction Made", "1) Implemented automatic form state reset after submission and added toast notification. 2) Added explicit empty state fallback component with 'Clear All Filters' recovery button.")

    # PART 6: SECURITY-AWARE DEVELOPMENT
    add_section_header("Part 6 — Security-Aware Development")
    add_key_value("Module 1 Misuse & Response", "Misuse: Submitting negative prices or XSS scripts via description field. | Response: Enforced client-side positive numeric bounds (minPrice > 0, quantity > 0), input sanitization, and Pydantic schema validation on backend endpoints.")
    add_key_value("Module 2 Misuse & Response", "Misuse: Placing unauthorized bids on expired/sold listings (biddingStatus !== 'OPEN'). | Response: State-level bidding lock disabling 'Place Bid' trigger for closed listings, paired with atomic server-side transaction verification.")

    # JAVASCRIPT INTERACTIONS LIST
    add_section_header("Summary of JavaScript Interactions Implemented")
    add_key_value("Interaction 1", "Real-Time Client-Side Validation & Stateful Form Reset (SubmitListing.jsx)")
    add_key_value("Interaction 2", "Multi-Attribute Search, Reactive Filter Chips, Array Sorting & Animated Listing Modal (BrowseListings.jsx)")

    file_name = "FairFarm_Activity3_Submission.docx"
    doc.save(file_name)
    print(f"Activity 3 Word Document created successfully: {file_name}")

def create_activity3_pdf():
    pdf_filename = "FairFarm_Activity3_Submission.pdf"
    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    PRIMARY = colors.HexColor("#555424") # Olive
    ACCENT = colors.HexColor("#853D25")  # Rust
    TEXT = colors.HexColor("#282828")

    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        textColor=PRIMARY,
        alignment=1, # Center
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        'SubTitleStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        textColor=ACCENT,
        alignment=1,
        spaceAfter=14
    )

    section_style = ParagraphStyle(
        'SectionStyle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        textColor=PRIMARY,
        spaceBefore=10,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        textColor=TEXT,
        leading=13,
        spaceAfter=4
    )

    bold_label = ParagraphStyle(
        'BoldLabel',
        parent=body_style,
        fontName='Helvetica-Bold'
    )

    story = []

    # Title
    story.append(Paragraph("DEVELOPMENT ACTIVITY 3 — PARALLEL MODULE DEVELOPMENT & INTEGRATION", title_style))
    story.append(Paragraph("FairFarm — Farmer Marketplace with Price Transparency | Secure Software Development", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=10))

    # Overview
    story.append(Paragraph("Submission Overview & Project Information", section_style))
    story.append(Paragraph("<b>Project Name:</b> FairFarm — Farmer Marketplace with Price Transparency", body_style))
    story.append(Paragraph("<b>Group Members:</b> Abdul-Wahab (2026-SSD-01), Ahmad-Fareed (2026-SSD-02)", body_style))
    story.append(Paragraph("<b>Public URL:</b> https://fair-farm-marketplace-p81iwnhzy-ogbrutal2825-8857s-projects.vercel.app/", body_style))
    story.append(Paragraph("<b>Module 1:</b> Submit Crop Listing (SubmitListing.jsx) — Developed by Abdul-Wahab", body_style))
    story.append(Paragraph("<b>Module 2:</b> Browse & Search Crop Listings (BrowseListings.jsx) — Developed by Ahmad-Fareed", body_style))

    # Part 1
    story.append(Paragraph("Part 1 — Selected Complementary Modules", section_style))
    story.append(Paragraph("<b>Module 1 (Submit Crop Listing):</b> Enables farmers to list agricultural produce with complete quality, quantity, price, and district metadata.", body_style))
    story.append(Paragraph("<b>Module 2 (Browse & Search Listings):</b> Enables buyers and officers to explore listings using real-time search, quality/status filters, sorting, and modal details.", body_style))

    # Part 2
    story.append(Paragraph("Part 2 — Member 1: Develop Module 1 (Abdul-Wahab)", section_style))
    story.append(Paragraph("<b>Responsible Member:</b> Abdul-Wahab", body_style))
    story.append(Paragraph("<b>Primary Function:</b> Produce submission form with live quality & price input fields.", body_style))
    story.append(Paragraph("<b>Validation Rules:</b> Mandatory fields, positive numeric checks (> 0), min 2 chars for variety.", body_style))
    story.append(Paragraph("<b>JS Interactions:</b> Real-time validation feedback, stateful session listing list, toast notification.", body_style))

    # Part 3
    story.append(Paragraph("Part 3 — Member 2: Develop Module 2 (Ahmad-Fareed)", section_style))
    story.append(Paragraph("<b>Responsible Member:</b> Ahmad-Fareed", body_style))
    story.append(Paragraph("<b>Primary Function:</b> Interactive produce catalog with multi-filter and inspection modal.", body_style))
    story.append(Paragraph("<b>Validation / Search Rules:</b> Multi-field search, status/quality filter chips, empty result recovery.", body_style))
    story.append(Paragraph("<b>JS Interactions:</b> Dynamic array filtering/sorting, filter chip removal, animated inspection modal.", body_style))

    # Part 4
    story.append(Paragraph("Part 4 — Integration Checklist", section_style))
    story.append(Paragraph("• Module 1 accessible at <code>/submit-listing</code> for authenticated Farmers.", body_style))
    story.append(Paragraph("• Module 2 accessible at <code>/browse-listings</code> for all users.", body_style))
    story.append(Paragraph("• Integrated into main Navbar and App router with unified brand theme.", body_style))

    # Part 5: Test Table
    story.append(Paragraph("Part 5 — Testing Matrix & Issues Discovered", section_style))
    table_data = [
        [Paragraph("<b>Test Case</b>", bold_label), Paragraph("<b>Expected Result</b>", bold_label), Paragraph("<b>Actual Result</b>", bold_label)],
        [Paragraph("Submit valid listing", body_style), Paragraph("Listing added, toast shown", body_style), Paragraph("Pass - Appended to sidebar list", body_style)],
        [Paragraph("Leave required empty", body_style), Paragraph("Form blocked, red error shown", body_style), Paragraph("Pass - Error messages displayed", body_style)],
        [Paragraph("Search & filter grid", body_style), Paragraph("Grid updates reactively", body_style), Paragraph("Pass - Cards filtered instantly", body_style)],
    ]
    t = Table(table_data, colWidths=[150, 180, 180])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#555424")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D1E2C4")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F4F6F0")]),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(t)
    story.append(Spacer(1, 6))

    story.append(Paragraph("<b>Module 1 Issue Discovered:</b> Inputs remained in form after submission. <i>Fix: Added automatic form state reset and success toast.</i>", body_style))
    story.append(Paragraph("<b>Module 2 Issue Discovered:</b> Empty search query broke layout. <i>Fix: Added dedicated empty state component with reset button.</i>", body_style))

    # Part 6
    story.append(Paragraph("Part 6 — Security-Aware Development", section_style))
    story.append(Paragraph("<b>Module 1 Misuse & Response:</b> Negative pricing / XSS injection -> Sanitization + positive numeric validation (> 0) + Pydantic backend check.", body_style))
    story.append(Paragraph("<b>Module 2 Misuse & Response:</b> Bidding on closed/sold items -> Bidding state lock disabling action + server transaction verification.", body_style))

    # Summary
    story.append(Paragraph("JavaScript Interactions Summary", section_style))
    story.append(Paragraph("1. Real-time form validation with inline errors and session listing state push.", body_style))
    story.append(Paragraph("2. Multi-criteria search, active filter chips, price sorting, and animated modal dialog.", body_style))

    doc.build(story)
    print(f"Activity 3 PDF created successfully: {pdf_filename}")

if __name__ == "__main__":
    create_activity3_docx()
    create_activity3_pdf()
