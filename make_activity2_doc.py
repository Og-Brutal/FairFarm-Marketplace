import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_activity2_doc():
    doc = docx.Document()

    # Standard Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colors
    PRIMARY_COLOR = RGBColor(85, 84, 36)   # Olive
    ACCENT_COLOR = RGBColor(133, 61, 37)   # Rust
    TEXT_COLOR = RGBColor(40, 40, 40)

    # Document Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.line_spacing = 1.2
    
    r_title = title_p.add_run("DEVELOPMENT ACTIVITY 2 - ROLE-BASED FUNCTIONAL PROTOTYPE")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(16)
    
    r_sub = sub_p.add_run("FairFarm - Farmer Marketplace with Price Transparency | Secure Software Development")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = ACCENT_COLOR

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    def add_key_value(label, val):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = TEXT_COLOR
        
        r_val = p.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)
        r_val.font.color.rgb = TEXT_COLOR

    # PART 1: SELECT TWO USER ROLES
    add_section_header("Part 1 - Selected User Roles")
    add_key_value("Role 1", "Farmer")
    add_key_value("Role 2", "Market Officer")

    # PART 2: DEFINE WHAT EACH ROLE CAN DO
    add_section_header("Part 2 - Role Responsibilities & Functions")
    
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Role 1 (Farmer) Functions:")
    r.bold = True
    r.font.color.rgb = ACCENT_COLOR
    
    farmer_funcs = [
        "1. Create, edit, and manage crop listings (specifying crop type, quantity, quality, and min price).",
        "2. Review incoming buyer bids and accept/confirm crop sales.",
        "3. View personal sales statistics, track active auctions, and verify KYC onboarding status."
    ]
    for fn in farmer_funcs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(fn)
        r.font.size = Pt(10)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Role 2 (Market Officer) Functions:")
    r.bold = True
    r.font.color.rgb = ACCENT_COLOR
    
    officer_funcs = [
        "1. Review and verify pending farmer/buyer KYC registrations (view unmasked CNIC and land record Fard numbers).",
        "2. Publish and update official market benchmark reference prices (e.g. Wheat, Rice, Cotton rates).",
        "3. Review suspicious bidding activity, monitor market audit logs, and resolve trading disputes."
    ]
    for fn in officer_funcs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(fn)
        r.font.size = Pt(10)

    # PART 3: CREATE DEMO USERS
    add_section_header("Part 3 - Demo Users")
    add_key_value("User 1 (Farmer)", "Name: Ahmed Khan | Email: farmer@fairfarm.pk | Password: password123 | Role: Farmer")
    add_key_value("User 2 (Market Officer)", "Name: Tariq Mahmood | Email: officer@fairfarm.pk | Password: password123 | Role: Market Officer")
    add_key_value("User 3 (Pending Farmer)", "Name: Mohammad Rashid | Email: rashid@fairfarm.pk | Password: password123 | Role: Farmer (Status: PENDING_VERIFICATION)")

    # PART 4: FUNCTIONAL LOGIN PAGE
    add_section_header("Part 4 - Functional Login Implementation")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("The login interface evaluates user credentials against demo role definitions or FastAPI backend authentication. Upon successful authentication, the system recognizes the user's role and automatically routes them to the appropriate dashboard: Farmers are directed to /marketplace while Market Officers are directed to /officer-dashboard.")
    r.font.size = Pt(10)

    # PART 5: ROLE-SPECIFIC DASHBOARDS
    add_section_header("Part 5 - Role-Specific Dashboards")
    add_key_value("Farmer Dashboard (/marketplace)", "Displays active listings, crop creation modal, personal bid statistics, live countdown timers, and an 'Attempt Restricted Action' button for security testing.")
    add_key_value("Market Officer Dashboard (/officer-dashboard)", "Displays Farmer KYC Verification queue (unmasked CNIC/Land records), Official Reference Price Publishing form, and Market Audit/Dispute resolution logs.")

    # PART 6: TEST ROLE DIFFERENCES
    add_section_header("Part 6 - Role Testing Results Table")
    
    test_rows = [
        ("Role 1 (Farmer) logs in", "Authenticated as Ahmed Khan; redirected to Farmer Marketplace (/marketplace)."),
        ("Role 2 (Market Officer) logs in", "Authenticated as Tariq Mahmood; redirected to Officer Dashboard (/officer-dashboard)."),
        ("Role 1 views dashboard", "Displays crop creation tools, personal sales metrics, and masked public PII data."),
        ("Role 2 views dashboard", "Displays unmasked CNIC/Land KYC verification queue, reference price publishing, and dispute logs.")
    ]
    for test, res in test_rows:
        add_key_value(test, res)

    # PART 7: PROTECT FUNCTION FROM WRONG ROLE
    add_section_header("Part 7 - Role-Restricted Function (RBAC Enforcement)")
    add_key_value("Function Selected", "Publish Official Market Reference Prices & Approve Farmer KYC Verification")
    add_key_value("Authorized Role", "Market Officer")
    add_key_value("Security Enforcement Logic", "IF user.role == 'Market Officer' THEN perform operation ELSE raise HTTP 403 Forbidden / display 'Access Denied: Only Market Officers are authorized to perform this operation.' END IF")

    # PART 8: SECURITY-AWARE DATA DISPLAY
    add_section_header("Part 8 - Security-Aware Data Display (PII Masking)")
    add_key_value("Selected Information", "Farmer CNIC Number (e.g. 35202-1234567-1) and Personal Phone Number (e.g. 0300-1234567)")
    add_key_value("Authorized Viewers", "Market Officers ONLY (Unmasked View: 35202-1234567-1). Other users/Farmers see Masked View (35202-*******-1) to prevent identity theft and privacy leaks.")

    # PART 9: LOGOUT IMPLEMENTATION
    add_section_header("Part 9 - Working Logout Implementation")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("Clicking 'Logout' clears active user role session state, wipes protected route tokens, resets navigation state back to 'Home', and redirects the browser immediately to the Login page. Direct URL navigation to protected dashboards after logout redirects to /login.")
    r.font.size = Pt(10)

    # PART 10: PROTOTYPE TESTING RESULTS
    add_section_header("Part 10 - System Prototype Test Execution")
    add_key_value("Test 1 - Correct Role", "Logged in as Ahmed Khan (Farmer) -> Recognized as Farmer, directed to Marketplace with crop management features.")
    add_key_value("Test 2 - Different Role", "Logged out & logged in as Tariq Mahmood (Market Officer) -> Recognized as Market Officer, directed to Officer Dashboard.")
    add_key_value("Test 3 - Restricted Function", "Attempted 'Publish Reference Price' while logged in as Farmer -> Access Denied (HTTP 403 Forbidden) alert modal triggered.")
    add_key_value("Test 4 - Logout", "Clicked Logout button -> Session cleared, nav badge removed, user returned to Login screen.")

    # PART 11: SECURITY DEVELOPMENT NOTE
    add_section_header("Part 11 - Security Development Note")
    add_key_value("Security Concept Applied", "Role-Based Access Control (RBAC) & PII Least Privilege Data Masking")
    add_key_value("Where Applied?", "FastAPI backend dependencies (require_role middleware) & React frontend route/action guards.")
    add_key_value("What Changed?", "Added role-based authorization headers, created role-specific dashboards, implemented CNIC masking functions, and built 403 Access Denied handling.")
    add_key_value("Unauthorized Operation Result", "Returns HTTP 403 Forbidden status code and renders an explicit 'Access Denied: Only Market Officers can perform this action' warning modal.")

    # FINAL SUBMISSION SUMMARY
    add_section_header("Submission Summary")
    add_key_value("Project Name", "FairFarm - Farmer Marketplace with Price Transparency")
    add_key_value("Group Members", "Abdul-Wahab, Ahmad-Fareed")
    add_key_value("Live/Local Project", "https://fair-farm-marketplace-p81iwnhzy-ogbrutal2825-8857s-projects.vercel.app/")
    add_key_value("Roles Implemented", "Farmer & Market Officer")
    add_key_value("Restricted Function", "Publish Official Market Reference Prices & Approve Farmer Verification")
    add_key_value("Security Concept Applied", "Role-Based Access Control (RBAC) & PII Data Masking (Least Privilege)")

    # Save Document
    file_name = "FairFarm_Activity2_Submission.docx"
    doc.save(file_name)
    print(f"Activity 2 Word Document created successfully: {file_name}")

if __name__ == "__main__":
    create_activity2_doc()
