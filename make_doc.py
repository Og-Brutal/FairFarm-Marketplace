import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_submission_doc():
    doc = docx.Document()

    # Standard Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Colors
    # Primary Olive: RGB(85, 84, 36)
    # Rust Accent: RGB(133, 61, 37)
    PRIMARY_COLOR = RGBColor(85, 84, 36)
    ACCENT_COLOR = RGBColor(133, 61, 37)
    TEXT_COLOR = RGBColor(40, 40, 40)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.line_spacing = 1.2
    
    run_title = title_p.add_run("FAIRFARM - PHASE 1 SUBMISSION")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(18)
    
    run_sub = sub_p.add_run("Farmer Marketplace with Price Transparency | Secure Software Development Project")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = ACCENT_COLOR

    def add_section_header(number_title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        run = p.add_run(number_title)
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    # 1. PROJECT INFORMATION
    add_section_header("1. Project Information")
    
    info_items = [
        ("Project Name", "FairFarm - Farmer Marketplace with Price Transparency"),
        ("Group Members", "Abdul-Wahab, Ahmad-Fareed"),
        ("Primary User Role", "Farmer (Primary Phase 1 role, with onboarding architecture for Buyer, Market Officer, Transport Provider, and System Administrator)")
    ]

    for label, val in info_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = TEXT_COLOR
        
        r_val = p.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)
        r_val.font.color.rgb = TEXT_COLOR

    # 2. LIVE WEBSITE
    add_section_header("2. Live Website")
    
    web_items = [
        ("Public URL", "https://fair-farm-marketplace-p81iwnhzy-ogbrutal2825-8857s-projects.vercel.app/"),
        ("GitHub Repository", "https://github.com/Og-Brutal/FairFarm-Marketplace.git")
    ]

    for label, val in web_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = TEXT_COLOR
        
        r_val = p.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)
        r_val.font.color.rgb = ACCENT_COLOR

    # 3. DEVELOPED PAGES
    add_section_header("3. Developed Pages")
    
    pages = [
        ("1. Home / Landing Page", "Features organic hero branding, live crop market price tickers, verified platform metrics (850+ Farmers, 2,400+ Listings, 100% Price Transparency), a 4-step workflow guide ('How It Works'), and an interactive Recharts price transparency chart."),
        ("2. Login & User Entry Page", "Split-screen authentication portal supporting dual modes: Farmer Sign-In and Public Farmer Registration with document upload preview and PENDING_VERIFICATION (KYC) onboarding status."),
        ("3. Farmer Marketplace Page", "Protected marketplace dashboard accessible exclusively to logged-in Farmers. Displays active listing stats, real-time crop search, status filters (Open, Ending Soon, Sold, Ended), live countdown timers, detailed crop modals, and a listing creation form.")
    ]

    for title, desc in pages:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_title = p.add_run(f"{title}: ")
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10.5)
        r_title.font.bold = True
        r_title.font.color.rgb = TEXT_COLOR
        
        r_desc = p.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = TEXT_COLOR

    # 4. JAVASCRIPT INTERACTIONS
    add_section_header("4. JavaScript / React Interactions")
    
    interactions = [
        ("1. Real-Time Animated Crop Countdown Timer", "Implemented using React useEffect and setInterval hooks to calculate remaining auction hours, minutes, and seconds every 1000ms. Automatically updates bidding status from 'OPEN' to 'ENDED' upon clock expiration."),
        ("2. Dynamic Price Transparency Chart & Search/Filter System", "Multi-crop tabbed price visualization powered by Recharts combined with dynamic client-side array filtering for crop name, district, quality grade, and bidding status."),
        ("3. Add Crop Listing Modal & Interactive State Management", "Modal overlay with form validation, district selectors, unit pickers, and dynamic React state mutations (setListings) that prepend new listings instantly to the marketplace grid with real-time recalculation of marketplace statistics.")
    ]

    for title, desc in interactions:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_title = p.add_run(f"{title}: ")
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10.5)
        r_title.font.bold = True
        r_title.font.color.rgb = TEXT_COLOR
        
        r_desc = p.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = TEXT_COLOR

    # 5. TESTING & DEBUGGING
    add_section_header("5. Testing & Debugging")
    
    testing_items = [
        ("Problem Discovered", "During crop card filtering and card hover state transitions, CSS 'transition-all' property rules conflicted with Framer Motion layout recalculations, causing grid layout stuttering and preventing the navbar active indicator line from updating smoothly during section scrolling."),
        ("Improvement Made", "Refactored navbar indicator state tracking using Framer Motion layoutId='navbar-indicator' with spring physics (stiffness: 380, damping: 30), removed conflicting CSS transition-all rules in favor of Framer Motion spring layout transitions, and implemented route protection for the marketplace using React Router DOM (<Navigate to='/login' replace />).")
    ]

    for label, val in testing_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = TEXT_COLOR
        
        r_val = p.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)
        r_val.font.color.rgb = TEXT_COLOR

    # Save document
    file_path = "FairFarm_Phase1_Submission.docx"
    doc.save(file_path)
    print(f"Document updated successfully: {file_path}")

if __name__ == "__main__":
    create_submission_doc()
